"""Populate the supplied dissertation template with verified project content."""
from pathlib import Path
from copy import deepcopy
import re, json, hashlib, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).parent
REFERENCE=ROOT/'tmp/thesis/template_reference.docx'
OUTPUT=ROOT/'reports/FireSense_AI_Dissertation_M_Janani_25MDT0227.docx'
TITLE='FireSense AI Fire and Smoke Detection Using YOLOv8'
doc=Document(REFERENCE)
sections=[deepcopy(s._sectPr) for s in doc.sections]
logo=deepcopy(doc.paragraphs[27]._p)
for item in logo.iter(qn('wp:docPr')):
    item.set('descr','VIT emblem and wordmark from the supplied dissertation template')
body=doc._element.body
for child in list(body):body.remove(child)

for style in ['Normal','Body Text','Caption','Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[style]
    st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(12)
    st.paragraph_format.space_before=Pt(0)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.widow_control=True
    st.paragraph_format.left_indent=Inches(0)
    st.paragraph_format.first_line_indent=Inches(0)
for name,size in [('Heading 1',18),('Heading 2',16),('Heading 3',14)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(8)
    st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.15
    st.paragraph_format.page_break_before=False
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.line_spacing=1.0
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER

def p(text='',style=None,align=None,size=None,bold=False,before=0,after=6,line=1.5):
    para=doc.add_paragraph(style=style)
    para.paragraph_format.space_before=Pt(before)
    para.paragraph_format.space_after=Pt(after)
    para.paragraph_format.line_spacing=line
    para.paragraph_format.widow_control=True
    para.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    r=para.add_run(text)
    if size:r.font.size=Pt(size)
    if bold:r.bold=True
    return para

def field(para,instruction,placeholder='1'):
    run=OxmlElement('w:r')
    for kind,text in [('begin',None),('instruction',instruction),('separate',None),('text',placeholder),('end',None)]:
        if kind=='instruction':
            el=OxmlElement('w:instrText');el.set(qn('xml:space'),'preserve');el.text=' '+text+' '
        elif kind=='text':el=OxmlElement('w:t');el.text=text
        else:el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
        run.append(el)
    para._p.append(run)

def boundary(index,body_start=False):
    sp=sections[index]
    typ=sp.find(qn('w:type'))
    if typ is None:typ=OxmlElement('w:type');sp.insert(0,typ)
    typ.set(qn('w:val'),'nextPage')
    num=sp.find(qn('w:pgNumType'))
    if num is not None:sp.remove(num)
    # Section properties govern the preceding content; numbering starts with section 7.
    if index==6:
        num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');num.set(qn('w:fmt'),'decimal');sp.append(num)
    if index<8:
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
        para._p.get_or_add_pPr().append(sp)
    else:body.append(sp)

def heading(text,level=1,page=False):
    para=doc.add_paragraph(text,style=f'Heading {level}')
    para.paragraph_format.page_break_before=page
    return para

def front_title(text):
    return p(text,align=WD_ALIGN_PARAGRAPH.CENTER,size=16,bold=True,before=18,after=24,line=1.2)

def table(rows):
    global tables
    tables+=1
    cap=p(style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER,line=1,after=6)
    cap.add_run('Table ');field(cap,'SEQ Table \\* ARABIC',str(tables));cap.add_run(' '+table_titles.get(tables,'Project evidence and configuration'))
    cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    width=6.24/len(rows[0])
    for col in t.columns:col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'A5A5A5');borders.append(el)
    props.append(borders)
    for i,row in enumerate(rows):
        tr=t.add_row();trpr=tr._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell,text in zip(tr.cells,row):
            cell.width=Inches(width)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shade)
            para=cell.paragraphs[0];para.paragraph_format.line_spacing=1.05
            para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
            para.paragraph_format.keep_with_next=i==0
            r=para.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=i==0
    p('',after=3,line=1).paragraph_format.keep_with_next=True

def figure(path,caption):
    global figures
    from PIL import Image
    target=(HERE/path).resolve()
    assert target.exists(),target
    with Image.open(target) as im:w,h=im.size
    width=min(6.2,4.8*w/h) if h>1.5*w else min(6.2,6.45*w/h)
    para=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=4,line=1)
    para.paragraph_format.keep_with_next=True
    r=para.add_run();shape=r.add_picture(str(target),width=Inches(width))
    for dp in shape._inline.iter(qn('wp:docPr')):dp.set('descr',caption)
    figures+=1
    cap=p(style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER,after=12,line=1.1)
    cap.add_run('Figure ');field(cap,'SEQ Figure \\* ARABIC',str(figures));cap.add_run(' '+caption)

tables=figures=0
table_titles={1:'Analysis of the starting state and completed work',2:'Recorded software and compute environment',3:'Dataset split and annotation inventory',4:'Repository responsibilities',5:'Original model training configuration',6:'Domain adaptation training configuration',7:'Completed implementation stages',8:'Model and deployment artifacts',9:'Current held-out test results',10:'Measured confidence operating points',11:'Model comparison with completed training budgets',12:'Adaptation epoch log',13:'Development-video observations',14:'Dataset and training evidence',15:'Metrics and checkpoint evidence',16:'Development and interface evidence'}

# Cover retains the original logo relationship and academic composition.
p(TITLE,style='Title',align=WD_ALIGN_PARAGRAPH.CENTER,size=19,bold=True,after=26,line=1.2)
p('Submitted in partial fulfillment of the requirements for the degree of',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,after=22,line=1.2).runs[0].italic=True
p('Master of Science',align=WD_ALIGN_PARAGRAPH.CENTER,size=22,bold=True,after=8,line=1.2)
p('In',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,after=8,line=1.2)
p('Data Science',align=WD_ALIGN_PARAGRAPH.CENTER,size=22,bold=True,after=22,line=1.2)
p('by',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,after=12,line=1.2).runs[0].italic=True
p('M Janani',align=WD_ALIGN_PARAGRAPH.CENTER,size=18,bold=True,after=6,line=1.2)
p('25MDT0227',align=WD_ALIGN_PARAGRAPH.CENTER,size=18,bold=True,after=22,line=1.2)
p('Under the guidance of',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,bold=True,after=14,line=1.2)
p('PRAVEEN T',align=WD_ALIGN_PARAGRAPH.CENTER,size=18,bold=True,after=6,line=1.2)
p('Assistant Professor Grade 2',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,after=6,line=1.2)
p('School of Advanced Sciences',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,after=6,line=1.2)
p('VIT Vellore',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,after=12,line=1.2)
body.append(logo)
lp=doc.paragraphs[-1];lp.alignment=WD_ALIGN_PARAGRAPH.CENTER
lp.paragraph_format.left_indent=Inches(0);lp.paragraph_format.first_line_indent=Inches(0)
lp.paragraph_format.space_before=Pt(0);lp.paragraph_format.space_after=Pt(10);lp.paragraph_format.line_spacing=1
p('October 2026',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,after=0,line=1.2)
boundary(0)

front_title('DECLARATION')
p('I, M Janani, registration number 25MDT0227, declare that the dissertation entitled "'+TITLE+'", submitted for the award of the degree of Master of Science in Data Science at VIT, is a record of the project work presented under the guidance of PRAVEEN T, Assistant Professor Grade 2, School of Advanced Sciences.')
p('The public dataset, software, publications and experimental evidence used in this dissertation are identified in the text and references. The reported measurements are drawn from the saved project results. External material has been acknowledged where it is used.')
p('I further declare that this dissertation has not been submitted for the award of another degree or diploma at this or another institution. This declaration is subject to the candidate\'s review and signature before submission.')
p('Place: Vellore',before=30)
p('Date: ____________________')
p('Signature of the candidate: ____________________',before=30)
p('M Janani',bold=True)
p('25MDT0227')
boundary(1)

front_title('CERTIFICATE')
p('This is to certify that the dissertation entitled "'+TITLE+'", submitted by M Janani, registration number 25MDT0227, towards the degree of Master of Science in Data Science at VIT, presents the project work described in this document under the guidance of PRAVEEN T, Assistant Professor Grade 2, School of Advanced Sciences.')
p('The certificate and examination details below are reserved for institutional review and approval. No signature, approval date or examination outcome is asserted in this unsigned copy.')
p('Guide signature: ____________________',before=26)
p('PRAVEEN T',bold=True)
p('Assistant Professor Grade 2')
p('School of Advanced Sciences')
p('Head of Department signature: ____________________',before=22)
p('Internal examiner: ____________________',before=16)
p('External examiner: ____________________',before=16)
p('Date of viva voce examination: ____________________',before=16)
boundary(2)

front_title('ACKNOWLEDGEMENT')
p('I thank my guide, PRAVEEN T, Assistant Professor Grade 2, School of Advanced Sciences, VIT, for the opportunity to present this project under his guidance.')
p('I acknowledge Bo Peng and Tae-Kook Kim for releasing the Home Fire Dataset, and the developers of Ultralytics and the open-source libraries used in the implementation. Their published work and software made it possible to build and evaluate the detector with a traceable technical basis.')
p('The project records both successful detections and the cases that required further work. Keeping those results visible has been an important part of preparing this dissertation.')
p('M Janani',align=WD_ALIGN_PARAGRAPH.RIGHT,before=36,bold=True)
p('25MDT0227',align=WD_ALIGN_PARAGRAPH.RIGHT)
boundary(3)

front_title('ABSTRACT')
p('FireSense AI is a local fire and smoke detection project built around YOLOv8. It provides dataset preparation, training, evaluation and a browser interface for camera frames, uploaded videos and still images. The project uses the publicly available Home Fire Dataset under CC BY-NC 4.0. An audit retained 6465 of its 6500 images, including 3873 training, 1297 validation and 1295 test images. The test split contains 1581 annotated objects across two classes.')
p('A YOLOv8 nano baseline completed 20 training epochs, while a small-model baseline completed eight under the available compute and time budget. The nano checkpoint was refined for four additional epochs and then adapted for three epochs using 600 close-up and portrait views derived only from the training split. This adaptation addressed a supplied video containing large flames under strong red lighting. The video influenced candidate selection and is therefore treated as development evidence.')
p('The selected model achieved 83.74% mAP@0.5 and 48.35% mAP@0.5:0.95 on the retained test split. At a minimum score of 0.50, precision was 93.04% and recall was 68.50%, with 1083 true positives, 81 false positives and 498 false negatives. Retained fire coverage improved from two to eight of eight fixed development frames. The interface displays raw model scores and supports annotated PNG and silent WebM downloads.')
p('The completed project is a runnable academic prototype with saved checkpoints, curves, prediction examples and commands. Its missed-object rate, uncertain source-level split independence and limited negative-video testing prevent a claim of dependable unattended fire warning. The dissertation documents those limits and proposes source-independent video evaluation before wider use.')
p('Keywords: fire detection, smoke detection, YOLOv8, object detection, local inference, domain adaptation.',size=12,line=1.2)
boundary(4)

front_title('TABLE OF CONTENTS')
para=p(line=1);field(para,'TOC \\o "1-2" \\h \\z \\u','Contents will update in Word')
doc.add_page_break()
front_title('LIST OF FIGURES')
para=p(line=1);field(para,'TOC \\h \\z \\c "Figure"','Figure list will update in Word')
doc.add_page_break()
front_title('LIST OF TABLES')
para=p(line=1);field(para,'TOC \\h \\z \\c "Table"','Table list will update in Word')
doc.add_page_break()
front_title('LIST OF ABBREVIATIONS')
for term,meaning in [('AI','Artificial intelligence'),('AP','Average precision'),('API','Application programming interface'),('CNN','Convolutional neural network'),('COCO','Common Objects in Context'),('CUDA','Compute Unified Device Architecture'),('DFL','Distribution focal loss'),('FP FN TP','False positive, false negative and true positive'),('FP32','32-bit floating point'),('GPU','Graphics processing unit'),('HSV','Hue saturation and value'),('HTTP','Hypertext Transfer Protocol'),('IoU','Intersection over union'),('mAP','Mean average precision'),('NMS','Non-maximum suppression'),('PNG','Portable Network Graphics'),('RGB','Red green and blue'),('SPPF','Spatial pyramid pooling fast'),('WCAG','Web Content Accessibility Guidelines'),('YOLO','You Only Look Once')]:
    p(term+'    '+meaning,line=1.1,after=7,size=12)
boundary(5)

lines=(HERE/'dissertation.md').read_text(encoding='utf-8').splitlines()
i=0;first=True
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        text=line[2:]
        if text.startswith('4 '):boundary(6);page=False
        elif text.startswith('10 '):boundary(7);page=False
        else:page=not first
        heading(text,1,page);first=False
    elif line.startswith('## '):heading(line[3:],2)
    elif line.startswith('@figure '):
        path,caption=line[8:].split(' | ',1);figure(path,caption)
    elif line.startswith('@code '):
        para=p(line[6:],align=WD_ALIGN_PARAGRAPH.LEFT,size=9,line=1.05,after=7)
        para.paragraph_format.keep_together=True
        for run in para.runs:run.font.name='Consolas'
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append([s.strip() for s in lines[i].strip().strip('|').split('|')]);i+=1
        table(rows);continue
    else:
        if line.startswith('The selected checkpoint hash is') or line.startswith('The source video\'s SHA256'):
            # Long hashes remain selectable while zero-width spaces permit safe wrapping.
            line=re.sub(r'\b[a-f0-9]{64}\b',lambda m:'\u200b'.join(m.group()[j:j+16] for j in range(0,64,16)),line)
        p(line)
    i+=1
boundary(8)

# Preserve the reference section geometry and border elements while repairing numbering.
for index,s in enumerate(doc.sections):
    s.different_first_page_header_footer=(index==0)
    s.first_page_footer.is_linked_to_previous=False
    for kind,footer in [('normal',s.footer),('first',s.first_page_footer)]:
        # The source stores PAGE inside content controls as well as direct paragraphs.
        for child in list(footer._element):footer._element.remove(child)
        para=footer.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing=1;para.paragraph_format.space_after=Pt(0)
        if kind=='normal':
            field(para,'PAGE','1')
            for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(10)

# Add a full-page visual audit selected from saved predictions and ground truth.
boundary(8)
heading('APPENDIX C HELD OUT TEST PREDICTION CASES',1,True)
p('Each page compares the dataset boxes with predictions from the selected YOLOv8 nano checkpoint. The score filter is 0.50, and a correct match requires the same class and an intersection over union of at least 0.50. Green outlines indicate fire; blue outlines indicate smoke. The first fifteen cases have no counted object errors under this rule. The remaining cases were selected from images with false positives or false negatives. They are examples for review, not a substitute for aggregate test metrics.',after=12)
case_manifest=json.loads((HERE/'figures/test_case_manifest.json').read_text())
for case in case_manifest:
    doc.add_page_break()
    p(f"Case {case['index']:02d}  {case['image']}",size=12,bold=True,align=WD_ALIGN_PARAGRAPH.LEFT,after=4,line=1.15)
    if case['kind']=='success':
        desc=f"This image contains {case['truth_count']} annotated object(s). The model returned {case['pred_count']} retained box(es), with {case['tp']} true positive(s), no false positives and no false negatives at the stated matching settings."
    else:
        desc=f"This image contains {case['truth_count']} annotated object(s). The model returned {case['pred_count']} retained box(es), with {case['tp']} true positive(s), {case['fp']} false positive(s) and {case['fn']} false negative(s). A retained box with inadequate overlap counts as an error under the matching rule."
    p(desc,size=12,after=8,line=1.25)
    figure(case['path'],f"Test case {case['index']:02d}, {case['image']}. Ground truth and current model predictions at score 0.50. Counts use class-aware matching at IoU 0.50.")

settings=doc.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
doc.core_properties.title=TITLE
doc.core_properties.subject='MSc Data Science project dissertation'
doc.core_properties.author='M Janani'
doc.core_properties.keywords='FireSense AI, YOLOv8, fire, smoke, 25MDT0227'
doc.core_properties.comments='Prepared from the supplied institutional template and saved project evidence.'
doc.save(OUTPUT)

assert len(doc.sections)==9
assert figures==48,(figures,tables)
assert tables==16,(figures,tables)
text='\n'.join(p.text for p in doc.paragraphs)
for unwanted in ['stock market','Tesla','Title of the Project','Name of the Student','November, 2025','—','–']:
    assert unwanted.lower() not in text.lower(),unwanted
with zipfile.ZipFile(REFERENCE) as a,zipfile.ZipFile(OUTPUT) as b:
    unchanged=[name for name in a.namelist() if name in b.namelist() and a.read(name)==b.read(name)]
audit=dict(reference_sha256=hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),output=str(OUTPUT),figures=figures,tables=tables,sections=9,unchanged_parts=unchanged,words=len((HERE/'dissertation.md').read_text().split()))
(ROOT/'tmp/thesis/build_checks.json').write_text(json.dumps(audit,indent=2))
print(json.dumps({k:v for k,v in audit.items() if k!='unchanged_parts'},indent=2))
